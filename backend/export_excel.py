"""Xuất bảng đáp án Excel theo hai bố cục mẫu: dọc và ngang.

Định dạng dữ liệu đầu vào:
{
  "code": "6001",
  "questions": [
    {"number": 1, "type": "multiple_choice", "answer": "A", "points": 0.25},
    {"number": 2, "type": "true_false", "answer": ["Đ", "S", "Đ", "S"], "points": 1},
    {"number": 3, "type": "short_answer", "answer": "2,5", "points": 0.5}
  ]
}
Không tự suy luận đáp án. Nếu answer là None hoặc chuỗi rỗng, ô đáp án để trống.
"""
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HEADERS_VERTICAL = ["Mã đề", "Câu hỏi", "Đáp án", "Điểm"]
TF_LABELS = ("a", "b", "c", "d")


def _text(value: Any) -> str:
    """Chuyển đáp án thành chuỗi Excel, không tạo đáp án thay cho dữ liệu thiếu."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "Đ" if value else "S"
    if isinstance(value, (list, tuple)):
        return ", ".join(str(item).strip() for item in value if item is not None and str(item).strip())
    return str(value).strip()


def _number(question: dict[str, Any], index: int) -> str:
    return str(question.get("number", index + 1))


def _type(question: dict[str, Any]) -> str:
    value = str(question.get("type", question.get("question_type", ""))).strip().lower()
    aliases = {
        "mcq": "multiple_choice", "multiple-choice": "multiple_choice", "choice": "multiple_choice",
        "đúng sai": "true_false", "dung_sai": "true_false", "true-false": "true_false", "tf": "true_false",
        "short-answer": "short_answer", "short answer": "short_answer", "trả lời ngắn": "short_answer",
    }
    return aliases.get(value, value)


def _answer(question: dict[str, Any]) -> str:
    value = question.get("answer")
    qtype = _type(question)
    if qtype == "true_false" and isinstance(value, (list, tuple)):
        # Giữ rõ đáp án từng tiểu ý, không gộp thành một chuỗi khó đối chiếu.
        return "; ".join(f"{label}) {_text(value[i])}" for i, label in enumerate(TF_LABELS) if i < len(value) and _text(value[i]))
    if qtype == "true_false" and isinstance(value, dict):
        return "; ".join(f"{label}) {_text(value[label] if label in value else value.get(label + ')'))}" for label in TF_LABELS if _text(value.get(label, value.get(label + ')'))))
    # MCQ phải là chữ cái A-D nếu nguồn có đáp án chữ cái.
    # Short answer giữ nguyên giá trị, không ép thành A-D.
    return _text(value)


def _rows_for_code(code: dict[str, Any]) -> list[dict[str, str]]:
    rows: list[dict[str, str]] = []
    questions = code.get("questions", [])
    if not isinstance(questions, list):
        raise ValueError("questions của mỗi mã đề phải là danh sách.")
    for index, question in enumerate(questions):
        if not isinstance(question, dict):
            raise ValueError("Mỗi câu trả lời phải là một đối tượng.")
        answer = _answer(question)
        qtype = _type(question)
        raw_answer = question.get("answer")
        if qtype in {"multiple_choice", "mcq", "choice"} and raw_answer is not None:
            candidate = _text(raw_answer).upper()
            if candidate and candidate not in {"A", "B", "C", "D"}:
                # Không biến dữ liệu không rõ thành chữ cái; vẫn xuất nguyên văn để rà soát.
                answer = _text(raw_answer)
        rows.append({
            "number": _number(question, index),
            "answer": answer,
            "points": question.get("points", ""),
        })
    return rows


def _style_sheet(ws, widths: dict[str, float]) -> None:
    header_fill = PatternFill("solid", fgColor="D9EAD3")
    thin = Side(style="thin", color="808080")
    for cell in ws[1]:
        cell.font = Font(name="Times New Roman", size=12, bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Times New Roman", size=12)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
    ws.row_dimensions[1].height = 28
    for col, width in widths.items():
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = ws.page_margins.right = 0.25
    ws.page_margins.top = ws.page_margins.bottom = 0.5
    ws.print_title_rows = "1:1"
    ws.auto_filter.ref = ws.dimensions if ws.max_row > 1 else "A1:A1"


def export_answer_key(codes: list[dict[str, Any]], path: str | Path, layout: str = "vertical") -> None:
    """Xuất đáp án theo mẫu dọc hoặc ngang.

    vertical/dọc/doc: Mã đề | Câu hỏi | Đáp án | Điểm.
    horizontal/ngang: Câu hỏi | mã đề 1 | mã đề 2 | ...
    Mỗi câu chiếm một dòng; đáp án đúng/sai được ghi rõ a) ...; b) ...;
    đáp án trả lời ngắn được giữ nguyên từ dữ liệu nguồn.
    """
    aliases = {"vertical": "vertical", "dọc": "vertical", "doc": "vertical", "horizontal": "horizontal", "ngang": "horizontal"}
    normalized = aliases.get(str(layout or "vertical").strip().lower())
    if normalized is None:
        raise ValueError("layout phải là vertical/dọc hoặc horizontal/ngang.")
    if not isinstance(codes, list) or not codes:
        raise ValueError("codes phải là danh sách có ít nhất một mã đề.")

    prepared = []
    seen_codes = set()
    for code in codes:
        if not isinstance(code, dict):
            raise ValueError("Mỗi mã đề phải là một đối tượng.")
        code_value = str(code.get("code", "")).strip()
        if not code_value:
            raise ValueError("Mỗi mã đề phải có trường code không rỗng.")
        if code_value in seen_codes:
            raise ValueError(f"Mã đề bị lặp: {code_value}")
        seen_codes.add(code_value)
        prepared.append((code_value, _rows_for_code(code)))

    wb = Workbook()
    ws = wb.active
    ws.title = "DapAn"
    if normalized == "vertical":
        ws.append(HEADERS_VERTICAL)
        for code, rows in prepared:
            for row in rows:
                ws.append([code, row["number"], row["answer"], row["points"]])
        _style_sheet(ws, {"A": 16, "B": 14, "C": 36, "D": 12})
    else:
        ws.append(["Câu hỏi"] + [code for code, _ in prepared])
        max_rows = max((len(rows) for _, rows in prepared), default=0)
        for index in range(max_rows):
            numbers = [rows[index]["number"] if index < len(rows) else "" for _, rows in prepared]
            # Các mã đề phải cùng cấu trúc câu hỏi; số câu được lấy từ mã đề đầu tiên.
            question_number = numbers[0] if numbers else str(index + 1)
            row = [question_number]
            for _, answers in prepared:
                row.append(answers[index]["answer"] if index < len(answers) else "")
            ws.append(row)
        widths = {get_column_letter(1): 16}
        widths.update({get_column_letter(i + 2): 28 for i in range(len(prepared))})
        _style_sheet(ws, widths)

    Path(path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    # Mở lại file để phát hiện lỗi ghi workbook trước khi API gửi cho người dùng.
    check = load_workbook(path, read_only=True, data_only=True)
    try:
        if check.active.max_row < 1:
            raise ValueError("Tệp Excel xuất ra không có dữ liệu.")
    finally:
        check.close()
