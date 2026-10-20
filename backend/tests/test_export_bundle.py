from io import BytesIO
from zipfile import ZipFile

from docx import Document
from fastapi.testclient import TestClient

from main import app

client = TestClient(app)
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def source_docx():
    doc = Document()
    doc.add_paragraph("PHẦN I. Trắc nghiệm nhiều phương án lựa chọn")
    doc.add_paragraph("Câu 1. Chọn đáp án đúng")
    for label, text in [("A", "Alpha"), ("B", "Beta"), ("C", "Gamma"), ("D", "Delta")]:
        doc.add_paragraph(f"{label}. {text}")
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_export_bundle_contains_docx_and_xlsx_and_uses_custom_code():
    response = client.post(
        "/api/exams/export-bundle",
        files={"file": ("source.docx", source_docx(), DOCX_MIME)},
        params={
            "exam_code": "KT12_A7",
            "header": '{"exam_title":"ĐỀ KIỂM TRA","subject":"TOÁN"}',
            "section_modes": '{"PI":"g0"}',
            "answers": '[{"number":1,"type":"multiple_choice","answer":"C","points":1}]',
            "answer_layout": "vertical",
            "seed": "42",
        },
    )
    assert response.status_code == 200, response.text
    assert response.headers["content-type"].startswith("application/zip")
    assert "KT12_A7" in response.headers.get("content-disposition", "")
    with ZipFile(BytesIO(response.content)) as archive:
        names = set(archive.namelist())
        assert "De-thi-ma-KT12_A7.docx" in names
        assert "Dap-an-ma-KT12_A7.xlsx" in names
        from openpyxl import load_workbook
        wb = load_workbook(BytesIO(archive.read("Dap-an-ma-KT12_A7.xlsx")), data_only=True)
        ws = wb.active
        assert ws.cell(1, 1).value == "Mã đề"
        assert ws.cell(2, 1).value == "KT12_A7"
        assert ws.cell(2, 3).value == "C"
        wb.close()


def test_export_bundle_rejects_invalid_custom_code():
    response = client.post(
        "/api/exams/export-bundle",
        files={"file": ("source.docx", source_docx(), DOCX_MIME)},
        params={
            "exam_code": "Mã đề 1",
            "answers": '[{"number":1,"type":"multiple_choice","answer":"A","points":1}]',
        },
    )
    assert response.status_code == 422
