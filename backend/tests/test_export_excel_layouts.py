from openpyxl import load_workbook
from export_excel import export_answer_key


def sample_codes():
    return [
        {"code": "6001", "questions": [
            {"number": 1, "answer": "D", "points": 1},
            {"number": 2, "answer": ["Đ", "S", "Đ", "S"], "points": 1},
        ]},
        {"code": "7002", "questions": [
            {"number": 1, "answer": "B", "points": 1},
            {"number": 2, "answer": ["S", "Đ", "S", "Đ"], "points": 1},
        ]},
    ]


def test_vertical_template(tmp_path):
    path = tmp_path / "vertical.xlsx"
    export_answer_key(sample_codes(), path, "vertical")
    ws = load_workbook(path, data_only=True).active
    assert [ws.cell(1, c).value for c in range(1, 5)] == ["Mã đề", "Câu hỏi", "Đáp án", "Điểm"]
    assert [ws.cell(2, c).value for c in range(1, 5)] == ["6001", "1", "D", 1]
    assert ws.cell(3, 3).value == "Đ, S, Đ, S"


def test_horizontal_template(tmp_path):
    path = tmp_path / "horizontal.xlsx"
    export_answer_key(sample_codes(), path, "horizontal")
    ws = load_workbook(path, data_only=True).active
    assert [ws.cell(1, c).value for c in range(1, 4)] == ["Câu hỏi", "6001", "7002"]
    assert [ws.cell(2, c).value for c in range(1, 4)] == ["1", "D", "B"]
    assert [ws.cell(3, c).value for c in range(1, 4)] == ["2", "Đ, S, Đ, S", "S, Đ, S, Đ"]
