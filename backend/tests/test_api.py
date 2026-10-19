from io import BytesIO
from pathlib import Path
import sys
import re

from docx import Document
from fastapi.testclient import TestClient

BACKEND_DIR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(BACKEND_DIR))
from main import app  # noqa: E402

client = TestClient(app)


def make_docx() -> bytes:
    doc = Document()
    doc.add_paragraph("Câu 1. Nước có công thức là gì?")
    doc.add_paragraph("A. H2O")
    doc.add_paragraph("B. CO2")
    doc.add_paragraph("C. O2")
    doc.add_paragraph("D. NaCl")
    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def test_health_endpoint():
    response = client.get("/api/health")
    assert response.status_code == 200
    assert response.json()["ok"] is True
    assert response.json()["service"] == "young-mix-pro-api"


def test_config_does_not_expose_secrets():
    response = client.get("/api/config")
    assert response.status_code == 200
    assert response.json()["features"]["pdf_import"] is False
    assert "GEMINI_API_KEY" not in response.text


def test_docx_parse_endpoint():
    response = client.post(
        "/api/exams/parse",
        files={"file": ("sample.docx", make_docx(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["question_count"] == 1
    assert len(body["questions"][0]["options"]) == 4
    assert body["review_required"] is True


def test_pdf_is_rejected_with_clear_message():
    response = client.post("/api/exams/parse", files={"file": ("sample.pdf", b"%PDF", "application/pdf")})
    assert response.status_code == 415
    assert "PDF/OCR" in response.json()["detail"]


def test_mix_rejects_empty_question_list():
    response = client.post("/api/exams/mix", json={"questions": [], "mode": "g3", "code_count": 2})
    assert response.status_code == 422


def make_two_question_docx() -> bytes:
    doc = Document()
    doc.add_paragraph("SỞ GIÁO DỤC VÀ ĐÀO TẠO")
    doc.add_paragraph("Câu 1. Câu thứ nhất?")
    doc.add_paragraph("A. Đáp án A")
    doc.add_paragraph("B. Đáp án B")
    doc.add_paragraph("Câu 2. Câu thứ hai?")
    doc.add_paragraph("A. Đáp án C")
    doc.add_paragraph("B. Đáp án D")
    doc.add_paragraph("HẾT")
    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def test_export_docx_preserves_source_and_reorders_questions():
    response = client.post(
        "/api/exams/export-docx?question_order=%5B%22Q-DOCX-P0-2%22%2C%22Q-DOCX-P0-1%22%5D",
        files={"file": ("sample.docx", make_two_question_docx(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert response.status_code == 200, response.text
    assert response.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    exported = Document(BytesIO(response.content))
    text = "\n".join(p.text for p in exported.paragraphs)
    assert "Câu 1. Câu thứ hai?" in text
    assert "Câu 2. Câu thứ nhất?" in text
    assert text.index("Câu 1. Câu thứ hai?") < text.index("Câu 2. Câu thứ nhất?")
    assert "SỞ GIÁO DỤC VÀ ĐÀO TẠO" in text
    assert "HẾT" in text


def test_export_docx_allows_exporting_selected_question_subset():
    response = client.post(
        "/api/exams/export-docx?question_order=%5B%22Q-DOCX-P0-2%22%5D",
        files={"file": ("sample.docx", make_two_question_docx(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert response.status_code == 200, response.text
    exported = Document(BytesIO(response.content))
    text = "\n".join(p.text for p in exported.paragraphs)
    assert "Câu 1. Câu thứ hai?" in text
    assert "Câu thứ nhất?" not in text
    assert "HẾT" in text



def test_parser_uses_unique_ids_when_question_numbers_repeat_between_sections():
    doc = Document()
    doc.add_paragraph("PHẦN I. Trắc nghiệm")
    doc.add_paragraph("Câu 1. Câu phần một?")
    doc.add_paragraph("A. Có")
    doc.add_paragraph("B. Không")
    doc.add_paragraph("PHẦN II. Đúng sai")
    doc.add_paragraph("Câu 1. Câu phần hai?")
    doc.add_paragraph("a) Ý thứ nhất")
    buffer = BytesIO()
    doc.save(buffer)
    response = client.post(
        "/api/exams/parse",
        files={"file": ("sections.docx", buffer.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert response.status_code == 200, response.text
    questions = response.json()["questions"]
    assert [q["id"] for q in questions] == ["Q-DOCX-PI-1", "Q-DOCX-PII-1"]


def test_export_docx_rejects_unknown_question_order():
    response = client.post(
        "/api/exams/export-docx?question_order=%5B%22Q-DOCX-999%22%5D",
        files={"file": ("sample.docx", make_two_question_docx(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert response.status_code == 422


def test_parser_returns_section_summary_for_immediate_mix_workflow():
    doc = Document()
    doc.add_paragraph("PHẦN I. Trắc nghiệm nhiều lựa chọn")
    doc.add_paragraph("Câu 1. Nước có công thức là gì?")
    doc.add_paragraph("A. H2O")
    doc.add_paragraph("B. CO2")
    doc.add_paragraph("PHẦN II. Trắc nghiệm đúng sai")
    doc.add_paragraph("Câu 1. Xét các phát biểu sau")
    doc.add_paragraph("a) Nước là H2O")
    buffer = BytesIO()
    doc.save(buffer)
    response = client.post(
        "/api/exams/parse",
        files={"file": ("parts.docx", buffer.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["question_count"] == 2
    assert [part["id"] for part in body["section_summary"]] == ["PI", "PII"]
    assert [part["question_count"] for part in body["section_summary"]] == [1, 1]
    assert body["questions"][1]["type"] == "true_false"


def test_parser_counts_inline_image_in_source_docx():
    import base64
    # Minimal 1x1 transparent PNG; avoid adding an image dependency to the backend.
    image_bytes = base64.b64decode("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+jS1sAAAAASUVORK5CYII=")
    doc = Document()
    doc.add_paragraph("PHẦN I. Trắc nghiệm")
    doc.add_paragraph("Câu 1. Quan sát hình sau:")
    doc.add_picture(BytesIO(image_bytes))
    doc.add_paragraph("A. Đúng")
    doc.add_paragraph("B. Sai")
    buffer = BytesIO()
    doc.save(buffer)
    response = client.post(
        "/api/exams/parse",
        files={"file": ("image.docx", buffer.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["image_count"] == 1
    assert body["questions"][0]["has_image"] is True


def test_export_docx_adds_official_header_and_removes_g_markers_and_answer_marks():
    from docx.shared import RGBColor
    doc = Document()
    doc.add_paragraph("<g3>")
    doc.add_paragraph("PHẦN I. Trắc nghiệm nhiều lựa chọn")
    doc.add_paragraph("Câu 1. Chọn phát biểu đúng?")
    correct = doc.add_paragraph()
    run = correct.add_run("A. Đáp án có đánh dấu")
    run.bold = True
    run.underline = True
    run.font.color.rgb = RGBColor(255, 0, 0)
    doc.add_paragraph("B. Phương án khác")
    doc.add_paragraph("ĐA:15")
    doc.add_paragraph("HẾT")
    doc.add_paragraph("ĐÁP ÁN")
    doc.add_paragraph("Câu 1: A")
    buf = BytesIO(); doc.save(buf)

    response = client.post(
        "/api/exams/export-docx",
        files={"file": ("sample.docx", buf.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        params={"header": '{"department":"SỞ GD&ĐT ĐỒNG NAI","school":"THPT THỐNG NHẤT B","exam_title":"KIỂM TRA CUỐI KỲ","grade":"12","school_year":"2026-2027","subject":"HÓA HỌC","duration":"45","exam_code":"007"}', "section_modes": '{"PI":"g0"}'},
    )
    assert response.status_code == 200, response.text
    assert response.headers["x-ymp-answer-markings-removed"] == "true"
    assert response.headers["x-ymp-answer-key-section-removed"] == "true"
    assert response.headers["x-ymp-parser-markers-removed"] == "true"
    out = Document(BytesIO(response.content))
    header_text = " ".join(cell.text for table in out.tables for row in table.rows for cell in row.cells)
    assert "SỞ GD&ĐT ĐỒNG NAI" in header_text
    assert "THPT THỐNG NHẤT B" in header_text
    assert "HÓA HỌC" in header_text
    all_text = "\n".join(p.text for p in out.paragraphs) + "\n" + header_text
    assert "<g3>" not in all_text
    assert "ĐA:15" not in all_text
    assert "Câu 1: A" not in all_text
    option = next(p for p in out.paragraphs if p.text.startswith("A. Đáp án có đánh dấu"))
    run = option.runs[0]
    assert not run.bold
    assert not run.underline
    assert run.font.color.rgb is None


def test_export_removes_A_value_lines_only_in_short_answer_section_and_adds_footer_page_field():
    from zipfile import ZipFile
    doc = Document()
    doc.add_paragraph("PHẦN I. Câu trắc nghiệm nhiều phương án lựa chọn")
    doc.add_paragraph("Câu 1. Chọn phương án đúng?")
    doc.add_paragraph("A. Phương án A")
    doc.add_paragraph("B. Phương án B")
    doc.add_paragraph("PHẦN III. Câu trắc nghiệm trả lời ngắn. Thí sinh trả lời từ câu 1 đến câu 2.")
    doc.add_paragraph("Câu 1. Tính giá trị biểu thức.")
    doc.add_paragraph("A. 10")
    doc.add_paragraph("Câu 2. Tính kết quả.")
    doc.add_paragraph("A. 25,5")
    buf = BytesIO(); doc.save(buf)
    response = client.post(
        "/api/exams/export-docx",
        files={"file": ("short-answer.docx", buf.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert response.status_code == 200, response.text
    out = Document(BytesIO(response.content))
    all_text = "\n".join(p.text for p in out.paragraphs)
    assert "A. Phương án A" in all_text
    assert "A. 10" not in all_text
    assert "A. 25,5" not in all_text
    assert out.sections[0].footer.paragraphs
    footer_text = "".join(p.text for p in out.sections[0].footer.paragraphs)
    assert "ĐỀ THI CHÍNH THỨC" not in footer_text
    assert "Mã đề: 001" in footer_text
    # Header carries live NUMPAGES and official designation; footer only has code + PAGE.
    with ZipFile(BytesIO(response.content)) as archive:
        footer_xml = archive.read("word/footer1.xml").decode("utf-8")
        document_xml = archive.read("word/document.xml").decode("utf-8")
    assert "PAGE" in footer_xml
    assert "NUMPAGES" in document_xml
    assert "ĐỀ THI CHÍNH THỨC" in document_xml
    assert "ĐỀ THI CHÍNH THỨC" not in footer_xml


def test_parser_reads_g3_marker_for_section_and_maps_part_four_to_essay():
    doc = Document()
    doc.add_paragraph("<g3>")
    doc.add_paragraph("PHẦN I. Trắc nghiệm nhiều lựa chọn")
    doc.add_paragraph("Câu 1. Chọn đáp án đúng")
    doc.add_paragraph("A. Một")
    doc.add_paragraph("B. Hai")
    doc.add_paragraph("PHẦN IV. Tự luận")
    doc.add_paragraph("Câu 1. Trình bày")
    buf = BytesIO(); doc.save(buf)
    result = client.post("/api/exams/parse", files={"file": ("markers.docx", buf.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
    assert result.status_code == 200, result.text
    body = result.json()
    assert body["questions"][0]["group"] == "g3"
    assert body["questions"][-1]["type"] == "essay"
    assert body["section_summary"][0]["group"] == "g3"


def test_export_g3_shuffles_labeled_option_paragraphs_and_relabels_them():
    doc = Document()
    doc.add_paragraph("<g3>")
    doc.add_paragraph("PHẦN I. Trắc nghiệm")
    doc.add_paragraph("Câu 1. Chọn")
    doc.add_paragraph("A. OPTION_ALPHA")
    doc.add_paragraph("B. OPTION_BETA")
    doc.add_paragraph("C. OPTION_GAMMA")
    doc.add_paragraph("D. OPTION_DELTA")
    buf = BytesIO(); doc.save(buf)
    response = client.post("/api/exams/export-docx", files={"file": ("shuffle.docx", buf.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}, params={"seed":"42"})
    assert response.status_code == 200, response.text
    out = Document(BytesIO(response.content))
    options = [p.text for p in out.paragraphs if re.match(r"^[A-D][.)]", p.text)]
    assert len(options) == 4
    assert [p[0] for p in options] == ["A", "B", "C", "D"]
    assert {p[3:] for p in options} == {"OPTION_ALPHA", "OPTION_BETA", "OPTION_GAMMA", "OPTION_DELTA"}
    assert response.headers["x-ymp-option-paragraphs-shuffled"] == "4"
