from openpyxl import load_workbook
from export_excel import export_answer_key


def sample_codes():
    return [
        {"code": "6001", "questions": [
            {"number": 1, "type": "multiple_choice", "answer": "A", "points": 0.25},
            {"number": 2, "type": "true_false", "answer": ["Đ", "S", "Đ", "S"], "points": 1},
            {"number": 3, "type": "short_answer", "answer": "2,5", "points": 0.5},
            {"number": 4, "type": "short_answer", "answer": None, "points": 0.5},
        ]},
        {"code": "7002", "questions": [
            {"number": 1, "type": "multiple_choice", "answer": "D", "points": 0.25},
            {"number": 2, "type": "true_false", "answer": ["S", "Đ", "S", "Đ"], "points": 1},
            {"number": 3, "type": "short_answer", "answer": "H2SO4", "points": 0.5},
            {"number": 4, "type": "short_answer", "answer": "", "points": 0.5},
        ]},
    ]


def test_vertical_outputs_mcq_tf_and_short_answers(tmp_path):
    path = tmp_path / "vertical.xlsx"
    export_answer_key(sample_codes(), path, "vertical")
    ws = load_workbook(path, data_only=True).active
    assert [ws.cell(1, c).value for c in range(1, 5)] == ["Mã đề", "Câu hỏi", "Đáp án", "Điểm"]
    assert [ws.cell(2, c).value for c in range(1, 5)] == ["6001", "1", "A", 0.25]
    assert ws.cell(3, 3).value == "a) Đ; b) S; c) Đ; d) S"
    assert ws.cell(4, 3).value == "2,5"
    assert ws.cell(5, 3).value is None or ws.cell(5, 3).value == ""


def test_horizontal_outputs_each_code_and_preserves_answers(tmp_path):
    path = tmp_path / "horizontal.xlsx"
    export_answer_key(sample_codes(), path, "horizontal")
    ws = load_workbook(path, data_only=True).active
    assert [ws.cell(1, c).value for c in range(1, 4)] == ["Câu hỏi", "6001", "7002"]
    assert [ws.cell(2, c).value for c in range(1, 4)] == ["1", "A", "D"]
    assert ws.cell(3, 2).value == "a) Đ; b) S; c) Đ; d) S"
    assert ws.cell(3, 3).value == "a) S; b) Đ; c) S; d) Đ"
    assert ws.cell(4, 2).value == "2,5"
    assert ws.cell(4, 3).value == "H2SO4"


def test_missing_answers_are_blank_and_not_guessed(tmp_path):
    path = tmp_path / "blank.xlsx"
    export_answer_key([{"code": "001", "questions": [{"number": 1, "type": "short_answer", "answer": None}]}], path)
    ws = load_workbook(path, data_only=True).active
    assert ws.cell(2, 3).value in (None, "")


def test_duplicate_codes_rejected(tmp_path):
    path = tmp_path / "duplicate.xlsx"
    data = [{"code": "001", "questions": []}, {"code": "001", "questions": []}]
    try:
        export_answer_key(data, path)
    except ValueError as exc:
        assert "lặp" in str(exc)
    else:
        raise AssertionError("Mã đề trùng phải bị từ chối")
